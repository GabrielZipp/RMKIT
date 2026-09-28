# ============================================
# RMKIT - MÓDULO DE LÓGICA (gerador.py)
# VERSÃO FINAL CORRIGIDA
# ============================================

import os
import sys
import tempfile
import shutil
import threading
import logging
from datetime import datetime
from pathlib import Path

import comtypes.client
import PyPDF2
import pythoncom
from docx import Document


# ---------- Logging ----------
LOG_PATH = Path(tempfile.gettempdir()) / "rmkit.log"
logging.basicConfig(
    filename=str(LOG_PATH),
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


class GeradorKit:

    # ---------- Helpers de data/texto ----------

    def data_por_extenso(self):
        meses = {
            1: 'janeiro', 2: 'fevereiro', 3: 'março', 4: 'abril',
            5: 'maio', 6: 'junho', 7: 'julho', 8: 'agosto',
            9: 'setembro', 10: 'outubro', 11: 'novembro', 12: 'dezembro'
        }
        hoje = datetime.now()
        return f"{hoje.day} de {meses[hoje.month]} de {hoje.year}"

    def gerar_frase_rg(self, rg_numero):
        if not rg_numero or not rg_numero.strip():
            return ""
        return f"portador(a) da cédula de identidade RG nº {rg_numero}"

    def montar_linha_honorarios_fixos(self, dados_cliente):
        incluir = self._bool_from_value(dados_cliente.get('{{INCLUIR_HON_FIXOS}}', True))
        valor = dados_cliente.get('{{HONORARIOS_FIXOS}}', '')
        if not incluir or not valor:
            return ""
        return f"a) Honorários fixos no valor de {valor}."

    def montar_linha_honorarios_exito(self, dados_cliente):
        incluir_exito = self._bool_from_value(dados_cliente.get('{{INCLUIR_HON_EXITO}}', True))
        valor = dados_cliente.get('{{HONORARIOS_EXITO}}', '')
        if not incluir_exito or not valor:
            return ""
        incluir_fixos = self._bool_from_value(dados_cliente.get('{{INCLUIR_HON_FIXOS}}', True))
        valor_fixos = dados_cliente.get('{{HONORARIOS_FIXOS}}', '')
        letra = 'b' if (incluir_fixos and valor_fixos) else 'a'
        return (f"{letra}) Honorários de êxito correspondentes a {valor} sobre os valores "
                f"líquidos efetivamente recebidos pelo(a) CONTRATANTE.")

    def _bool_from_value(self, value):
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            return value.lower() in ('true', '1', 'yes')
        return bool(value)

    # ---------- Substituição de placeholders (PRESERVA FORMATAÇÃO) ----------

    def substituir_no_paragrafo(self, paragrafo, substituicoes):
        """Substitui placeholders preservando a formatação de cada run.

        Aplica em loop porque uma substituição pode expor outro placeholder
        (ex.: se um valor contiver outro placeholder)."""
        if not paragrafo.runs:
            return
        for _ in range(10):
            if not self._aplicar_substituicoes(paragrafo, substituicoes):
                break

    def _aplicar_substituicoes(self, paragrafo, substituicoes):
        """Uma passada. Retorna True se algo foi alterado."""
        runs = list(paragrafo.runs)
        if not runs:
            return False

        textos = [r.text for r in runs]
        alterou = False

        for placeholder, valor in substituicoes.items():
            valor = "" if valor is None else str(valor)

            while True:
                texto_completo = "".join(textos)
                idx = texto_completo.find(placeholder)
                if idx == -1:
                    break

                fim = idx + len(placeholder)

                # Localiza qual run contém o início e qual contém o fim
                acc = 0
                inicio_run = fim_run = None
                for i, t in enumerate(textos):
                    if inicio_run is None and acc + len(t) > idx:
                        inicio_run = i
                    if acc + len(t) >= fim:
                        fim_run = i
                        break
                    acc += len(t)

                if inicio_run is None or fim_run is None:
                    break  # segurança

                # Offset do início do placeholder dentro do run inicial
                acc_ini = sum(len(textos[i]) for i in range(inicio_run))
                offset_ini = idx - acc_ini

                # Offset do fim do placeholder dentro do run final
                acc_fim = sum(len(textos[i]) for i in range(fim_run))
                offset_fim = fim - acc_fim

                if inicio_run == fim_run:
                    # Placeholder inteiro em um único run — não perde formatação
                    textos[inicio_run] = (
                        textos[inicio_run][:offset_ini]
                        + valor
                        + textos[inicio_run][offset_fim:]
                    )
                else:
                    # Placeholder dividido entre runs:
                    # - o valor entra no run inicial, herdando a formatação dele
                    # - runs intermediários são esvaziados (mas preservam formatação)
                    # - o run final mantém só o que vem DEPOIS do placeholder
                    prefixo = textos[inicio_run][:offset_ini]
                    sufixo = textos[fim_run][offset_fim:]

                    textos[inicio_run] = prefixo + valor
                    for i in range(inicio_run + 1, fim_run):
                        textos[i] = ""
                    textos[fim_run] = sufixo

                alterou = True

        if alterou:
            for i, r in enumerate(runs):
                if r.text != textos[i]:
                    r.text = textos[i]

        return alterou

    def substituir_no_tabela(self, tabela, substituicoes):
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    self.substituir_no_paragrafo(paragrafo, substituicoes)

    def substituir_placeholders_docx(self, docx_path, substituicoes):
        doc = Document(str(docx_path))
        for paragrafo in doc.paragraphs:
            self.substituir_no_paragrafo(paragrafo, substituicoes)
        for tabela in doc.tables:
            self.substituir_no_tabela(tabela, substituicoes)
        for secao in doc.sections:
            if secao.header:
                for p in secao.header.paragraphs:
                    self.substituir_no_paragrafo(p, substituicoes)
                for t in secao.header.tables:
                    self.substituir_no_tabela(t, substituicoes)
            if secao.footer:
                for p in secao.footer.paragraphs:
                    self.substituir_no_paragrafo(p, substituicoes)
                for t in secao.footer.tables:
                    self.substituir_no_tabela(t, substituicoes)
        doc.save(str(docx_path))

    # ---------- Conversão Word → PDF ----------

    def _matar_word_orfao(self):
        """Fallback agressivo quando o Word trava."""
        try:
            if sys.platform == 'win32':
                os.system('taskkill /F /IM WINWORD.EXE /T >nul 2>&1')
        except Exception:
            pass

    def gerar_pdf_individual(self, template_path, dados_cliente, output_pdf, timeout=60):
        substituicoes = dados_cliente.copy()
        substituicoes['{{DATA_EXTENSO}}'] = self.data_por_extenso()
        substituicoes['{{FRASE_RG}}'] = self.gerar_frase_rg(dados_cliente.get('{{RG}}', ''))
        substituicoes['{{FRASE_HONORARIOS_FIXOS}}'] = self.montar_linha_honorarios_fixos(dados_cliente)
        substituicoes['{{FRASE_HONORARIOS_EXITO}}'] = self.montar_linha_honorarios_exito(dados_cliente)
        for chave, valor in substituicoes.items():
            if valor is None:
                substituicoes[chave] = ""

        temp_dir = Path(tempfile.mkdtemp())
        temp_docx = temp_dir / "temp.docx"
        shutil.copy2(str(template_path), str(temp_docx))

        resultado = {"ok": False, "erro": None}
        holder = {"doc": None, "word": None}

        try:
            self.substituir_placeholders_docx(temp_docx, substituicoes)

            def converter():
                try:
                    pythoncom.CoInitialize()
                    word = comtypes.client.CreateObject("Word.Application")
                    word.Visible = False
                    word.DisplayAlerts = False
                    holder["word"] = word

                    doc = word.Documents.Open(str(temp_docx.absolute()))
                    holder["doc"] = doc

                    doc.SaveAs(str(Path(output_pdf).absolute()), FileFormat=17)
                    doc.Close(SaveChanges=False)
                    holder["doc"] = None

                    word.Quit()
                    holder["word"] = None
                    resultado["ok"] = True
                except Exception as e:
                    resultado["erro"] = str(e)
                    logger.exception("Erro na conversão Word->PDF")
                finally:
                    try:
                        pythoncom.CoUninitialize()
                    except Exception:
                        pass

            thread = threading.Thread(target=converter, daemon=True)
            thread.start()
            thread.join(timeout=timeout)

            if thread.is_alive():
                logger.error("Timeout na conversão; tentando liberar Word.")
                for obj in ("doc", "word"):
                    try:
                        if holder[obj] is not None:
                            if obj == "doc":
                                holder[obj].Close(SaveChanges=False)
                            else:
                                holder[obj].Quit()
                            holder[obj] = None
                    except Exception:
                        pass
                self._matar_word_orfao()
                raise Exception(
                    f"Tempo limite excedido ({timeout}s). Verifique se o Word não está "
                    f"com pop-up ou travado."
                )
            if resultado["erro"]:
                raise Exception(resultado["erro"])
            if not resultado["ok"]:
                raise Exception("Falha na conversão para PDF")
            return output_pdf
        finally:
            try:
                shutil.rmtree(temp_dir, ignore_errors=True)
            except Exception:
                pass

    # ---------- Kit completo ----------

    def gerar_kit(self, template_proc, template_cont, dados_cliente, output_pdf):
        if not os.path.exists(template_proc):
            raise FileNotFoundError(f"Template procuração não encontrado: {template_proc}")
        if not os.path.exists(template_cont):
            raise FileNotFoundError(f"Template contrato não encontrado: {template_cont}")

        kit_temp_dir = Path(tempfile.mkdtemp())
        pdf_proc = kit_temp_dir / "procuracao.pdf"
        pdf_cont = kit_temp_dir / "contrato.pdf"

        try:
            self.gerar_pdf_individual(template_proc, dados_cliente, str(pdf_proc))
            self.gerar_pdf_individual(template_cont, dados_cliente, str(pdf_cont))

            merger = PyPDF2.PdfMerger()
            try:
                merger.append(str(pdf_proc))
                merger.append(str(pdf_cont))
                merger.write(str(output_pdf))
            finally:
                merger.close()
            return output_pdf
        finally:
            shutil.rmtree(kit_temp_dir, ignore_errors=True)

    def validar_dados(self, dados):
        obrigatorios = ['{{NOME}}', '{{CPF}}']
        faltando = [campo for campo in obrigatorios if not dados.get(campo)]
        return len(faltando) == 0, faltando